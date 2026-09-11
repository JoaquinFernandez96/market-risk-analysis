"""Unit tests for VaR supporting schedules and Excel audit export."""
import io
import pytest
import openpyxl
from src.portfolio.models import create_default_portfolio
from src.risk.var_engine import VaREngine
from src.risk.schedules import (
    generate_market_value_evolution,
    generate_returns_pnl_schedule,
    generate_parametric_audit_schedules,
    generate_tail_loss_audit_schedule,
    build_var_supporting_excel
)


@pytest.fixture
def test_setup():
    portfolio = create_default_portfolio()
    var_engine = VaREngine(portfolio)
    ccl_rate = 1582.60
    return portfolio, var_engine, ccl_rate


def test_market_value_evolution(test_setup):
    portfolio, var_engine, ccl_rate = test_setup
    df_mv = generate_market_value_evolution(var_engine.base_returns_df, portfolio, ccl_rate)

    assert len(df_mv) == 504
    assert "Portfolio_NAV_USD" in df_mv.columns
    assert "Portfolio_NAV_ARS" in df_mv.columns

    # Final row must match current portfolio NAV
    final_usd = df_mv["Portfolio_NAV_USD"].iloc[-1]
    assert pytest.approx(final_usd, rel=1e-3) == portfolio.total_value

    # ARS column must match USD * ccl_rate
    final_ars = df_mv["Portfolio_NAV_ARS"].iloc[-1]
    assert pytest.approx(final_ars, rel=1e-3) == (portfolio.total_value * ccl_rate)

    # Check that individual asset columns exist in dual currencies for both UnitPrice and TotalVal
    for pos in portfolio.positions:
        assert f"{pos.ticker}_TotalVal_USD" in df_mv.columns
        assert f"{pos.ticker}_TotalVal_ARS" in df_mv.columns
        assert f"{pos.ticker}_UnitPrice_USD" in df_mv.columns
        assert f"{pos.ticker}_UnitPrice_ARS" in df_mv.columns
        assert pytest.approx(df_mv[f"{pos.ticker}_TotalVal_USD"].iloc[-1], rel=1e-3) == pos.market_value
        if pos.quantity > 0:
            expected_unit_p = pos.market_value / pos.quantity
            assert pytest.approx(df_mv[f"{pos.ticker}_UnitPrice_USD"].iloc[-1], rel=1e-2) == expected_unit_p


def test_returns_pnl_schedule(test_setup):
    portfolio, var_engine, ccl_rate = test_setup
    df_mv = generate_market_value_evolution(var_engine.base_returns_df, portfolio, ccl_rate)
    df_pnl = generate_returns_pnl_schedule(var_engine.base_returns_df, df_mv, portfolio, ccl_rate)

    assert len(df_pnl) == 504
    assert "Portfolio_Daily_PnL_USD" in df_pnl.columns
    assert "Portfolio_Daily_PnL_ARS" in df_pnl.columns
    assert "Portfolio_Return(%)" in df_pnl.columns


def test_parametric_audit_schedules(test_setup):
    portfolio, var_engine, ccl_rate = test_setup
    cov_df, corr_df, comp_df = generate_parametric_audit_schedules(
        var_engine, "base", ccl_rate, confidence_level=0.95, horizon_days=1
    )

    assert cov_df.shape[0] == len(portfolio.positions)
    assert corr_df.shape[0] == len(portfolio.positions)
    assert len(comp_df) == len(portfolio.positions)
    assert "Component VaR (ARS)" in comp_df.columns


def test_build_var_supporting_excel(test_setup):
    portfolio, var_engine, ccl_rate = test_setup
    excel_bytes = build_var_supporting_excel(
        portfolio, var_engine, ccl_rate, confidence_level=0.95, horizon_days=1
    )

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 5000

    # Verify openpyxl can load and read all 6 sheets
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    expected_sheets = [
        "Base_Market_Values",
        "Base_Returns_PnL",
        "Stressed_Market_Values",
        "Stressed_Returns_PnL",
        "Parametric_Risk_Matrix",
        "VaR_CVaR_Tail_Audit"
    ]
    for s in expected_sheets:
        assert s in wb.sheetnames
        assert wb[s].max_row > 5
